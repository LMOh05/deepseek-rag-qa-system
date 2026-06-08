"""
文档解析模块 —— 支持 PDF、Word、TXT，统一输出 LangChain Document 列表。
"""

from pathlib import Path
from typing import List
from langchain_core.documents import Document


def load_pdf(file_path: Path) -> List[Document]:
    """解析 PDF，每页作为一个 Document"""
    from PyPDF2 import PdfReader

    reader = PdfReader(str(file_path))
    docs = []
    for i, page in enumerate(reader.pages):
        text = page.extract_text()
        if text and text.strip():
            docs.append(Document(
                page_content=text.strip(),
                metadata={
                    "source": str(file_path),
                    "page": i + 1,
                    "type": "pdf",
                    "filename": file_path.name,
                }
            ))
    return docs


def load_docx(file_path: Path) -> List[Document]:
    """解析 Word 文档，按段落分组"""
    from docx import Document as DocxDocument

    doc = DocxDocument(str(file_path))
    paragraphs = []
    current = ""
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            if current:
                paragraphs.append(current)
                current = ""
            continue
        # 按标题样式分段
        if para.style.name.startswith("Heading") and current:
            paragraphs.append(current)
            current = text
        else:
            current = (current + "\n" + text) if current else text

    if current:
        paragraphs.append(current)

    return [
        Document(
            page_content=p,
            metadata={
                "source": str(file_path),
                "type": "docx",
                "filename": file_path.name,
                "chunk_index": i,
            }
        )
        for i, p in enumerate(paragraphs) if p.strip()
    ]


def load_txt(file_path: Path) -> List[Document]:
    """解析 TXT 文本，按空行自然分段，合并过小段落避免碎片化。"""
    text = file_path.read_text(encoding="utf-8")
    # 按空行分段
    raw_paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]

    if not raw_paragraphs:
        raw_paragraphs = [text.strip()]

    # 合并过小的连续段落，避免产生无意义的碎片（如孤立标题）
    MIN_PARAGRAPH_CHARS = 100
    merged = []
    buffer = ""
    for p in raw_paragraphs:
        if len(buffer) + len(p) <= MIN_PARAGRAPH_CHARS:
            buffer = (buffer + "\n" + p).strip() if buffer else p
        else:
            if buffer:
                merged.append(buffer)
            buffer = p
    if buffer:
        # 最后一个 buffer 如果还太小，合并到前一个段落
        if merged and len(buffer) < MIN_PARAGRAPH_CHARS:
            merged[-1] = merged[-1] + "\n" + buffer
        else:
            merged.append(buffer)

    return [
        Document(
            page_content=p,
            metadata={
                "source": str(file_path),
                "type": "txt",
                "filename": file_path.name,
                "chunk_index": i,
            }
        )
        for i, p in enumerate(merged) if p.strip()
    ]


LOADER_MAP = {
    ".pdf": load_pdf,
    ".docx": load_docx,
    ".doc": load_docx,
    ".txt": load_txt,
    ".md": load_txt,
}


def load_documents(data_dir: str = None) -> List[Document]:
    """
    递归扫描目录，加载所有支持的文档。

    Args:
        data_dir: 文档目录路径，默认使用 config.DATA_DIR

    Returns:
        List[Document]: 所有文档的 LangChain Document 列表
    """
    from src.config import config

    if data_dir is None:
        data_dir = config.DATA_DIR

    data_path = Path(data_dir)
    if not data_path.exists():
        raise FileNotFoundError(f"文档目录不存在: {data_dir}")

    all_docs = []
    stats = {}  # 按类型统计

    for file_path in data_path.rglob("*"):
        if not file_path.is_file():
            continue

        ext = file_path.suffix.lower()
        if ext not in LOADER_MAP:
            continue

        try:
            docs = LOADER_MAP[ext](file_path)
            all_docs.extend(docs)
            stats[ext] = stats.get(ext, 0) + len(docs)
            print(f"  OK [{ext}] {file_path.name} -> {len(docs)} paragraphs")
        except Exception as e:
            print(f"  FAIL [{ext}] {file_path.name}: {e}")

    print(f"\n[Docs] Loaded: {len(all_docs)} paragraphs")
    for ext, count in sorted(stats.items()):
        print(f"   {ext}: {count}")
    return all_docs


if __name__ == "__main__":
    docs = load_documents()
    for doc in docs[:3]:
        print(f"\n--- {doc.metadata['filename']} ---")
        print(doc.page_content[:300])
